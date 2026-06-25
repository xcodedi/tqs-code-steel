import os
import re
from collections import defaultdict
from openpyxl import load_workbook

os.system('cls' if os.name == 'nt' else 'clear')

# ======= CONFIGURAÇÕES =======
PASTA = r"c:\caminho\da\pasta" #caminho da pasta aonde estão os lst
BITOLAS_VALIDAS = {5, 6.3, 8, 10, 12.5, 16, 20,22,25,32,40,50}
# =============================

# Regex para ler linhas válidas de aço
LINHA_RESUMO = re.compile(
    r'^\s*(50A|60A)\s+(\d+(?:[.,]\d+)?)\s+(\d+(?:[.,]\d+)?)\s+(\d+(?:[.,]\d+)?)\s*$'
)

# Regex para extrair pavimento e elemento do nome do arquivo
PADRAO_NOME = re.compile(
    r'(?i)\d{3}[-_](FUN|PIL|TER|TÉR|HAL|G1_|ELE|COB|TEL|MAQ|CAI|CIN|PAT|ESP|SUB1|GAR|GAR2|LAZ|RAM|SUB2|G2_|TIP1|SAL|TIP5|TIP69|TIP1015|TIP1619|TIP2025|TIP2627|DUP|DUPS|MOT|CAL|PAT2|TIP)[-_](LAJ|VIG|ESC|FOR|OUT|FUN|PIL)[-_]\d{3}'
)

# Dicionários de tradução
MAPA_PAVIMENTO = {
    # infraestrutura
    "FUN": "Fundação",
    "PIL": "Pilares",

    # Subsolos
    "SUB2": "Subsolo-2",
    "SUB1": "Subsolo-1",

    # Acessos
    "RAM": "Rampa",
    "CAL": "Calçada",
    "TER": "Térreo",
    "TÉR": "Térreo",
    "HAL": "Hall",

    # Garagens
    "GAR": "Garagem-1",
    "GAR1": "G1",
    "G1_": "GAR1",

    "GAR2": "Garagem-2",
    "G2_": "GAR2",

    "ELE": "Elevado-lazer",

    # Pavimentos
    "LAZ": "Lazer",
    "TIP": "TIPO",
    "TIP1": "Tipo-1",
    "TIP5": "PAV 5",
    "TIP69": "Tipo-6-9",
    "TIP7A14": "Tipo-7-14",
    "TIP1015": "Tipo-10-15",
    "TIP15A23": "Tipo-15-23",
    "TIP1619": "Tipo-16-19",
    "TIP2025": "Tipo-20-25",
    "TIP2627": "Tipo-26-27",

    # Duplex
    "DUP": "Duplex-inferior",
    "PAT": "Patamar",
    "PAT2": "PATAMAR-SUPERIOR",
    "DUPS": "Duplex-superior",

    # Áreas superiores
    "SAL": "Salão-festas",
    "ESP": "Espaço Técnico",
    "COB": "Cobertura",
    "TEL": "Telhado",

    # Técnicos
    "MAQ": "Casa de Máquinas",
    "MOT": "MOTOR",
    "CAI": "Caixa",
    "CIN": "Cinta",
}

MAPA_ELEMENTO = {
    "LAJ": "Lajes",
    "VIG": "Vigas",
    "ESC": "Escadas",
    "FOR": "Formas",
    "OUT": "Outros Elementos",
    "PIL": "Pilares",
    "FUN": "Blocos"
}

#Converte texto numérico para float, aceita vírgula e ponto
def parse_num(s: str) -> float | None:
    
    t = s.strip()
    if "," in t and "." in t:
        # Formato tipo 1.685,00 -> 1685.00
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None

#Processa um arquivo .LST e retorna a soma dos pesos por bitola
def processar_lst(caminho_arquivo: str):

    for enc in ("cp1252", "latin-1", "utf-8"):
        try:
            linhas = open(
                caminho_arquivo,
                "r",
                encoding=enc,
                errors="ignore"
            ).read().splitlines()
            break
        except Exception:
            linhas = None

    soma_bitola = defaultdict(float)
    soma_aco = defaultdict(float)

    for linha in linhas:

        if "peso total" in linha.lower():
            continue

        m = LINHA_RESUMO.match(linha)

        if not m:
            continue

        aco, bit_s, _, peso_s = m.groups()

        bit = parse_num(bit_s)
        peso = parse_num(peso_s)

        if bit is None or peso is None or bit not in BITOLAS_VALIDAS:
            continue

        soma_bitola[bit] += peso
        soma_aco[aco] += peso

    return soma_bitola, soma_aco


# VARREDURA 
arquivos = [f for f in os.listdir(PASTA) if f.lower().endswith(".lst")]
if not arquivos:
    print("Nenhum arquivo .LST encontrado na pasta. Verifique PASTA.")
    exit()

# Estrutura: {pavimento: {elemento: {bitola: peso}}}
dados = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
total_geral_bitola = defaultdict(float)
total_geral_aco = defaultdict(float)
for nome in arquivos:
    caminho = os.path.join(PASTA, nome)
    soma_folha_bitola, soma_folha_aco = processar_lst(caminho)

    m = PADRAO_NOME.search(nome)
    if not m:
        print(f" Nome do arquivo fora do padrão: {nome}")
        continue

    pav_raw, elem_raw = m.groups()

    pavimento = MAPA_PAVIMENTO.get(pav_raw.upper())
    elemento = MAPA_ELEMENTO.get(elem_raw.upper())

    if not pavimento:
        print(f" Pavimento '{pav_raw}' não mapeado (arquivo: {nome})")
        pavimento = pav_raw

    if not elemento:
        print(f" Elemento '{elem_raw}' não mapeado (arquivo: {nome})")
        elemento = elem_raw

    # TEM QUE FICAR DENTRO DO LOOP

    for bit, peso in soma_folha_bitola.items():
        dados[pavimento][elemento][bit] += peso

    for bit, peso in soma_folha_bitola.items():
        total_geral_bitola[bit] += peso

    for aco, peso in soma_folha_aco.items():
        total_geral_aco[aco] += peso

# EXIBE RESULTADOS
print("\n TOTAL DE AÇO POR PAVIMENTO, ELEMENTO E BITOLA:\n")

total_geral = 0

for pav, elementos in dados.items():
    print(f" {pav.upper()}:")
    total_pav = 0
    for elem, bitolas in elementos.items():
        total_elem = sum(bitolas.values())
        total_pav += total_elem
        print(f"  - {elem}: total {total_elem:.2f} kgf")
        for bit in sorted(bitolas):
            print(f"      Bitola {bit:g} mm: {bitolas[bit]:.2f} kgf")
        print()
    total_geral += total_pav
    print(f"  Total do pavimento {pav}: {total_pav:.2f} kgf\n")

print("\n" + "=" * 80)
print(f"\n TOTAL GERAL DE AÇO: {total_geral:.2f} kgf\n")
print("=" * 80)
print("\nTOTAL GERAL POR BITOLA:")

for bit in sorted(total_geral_bitola):
    print(
        f"  Bitola {bit:g} mm: "
        f"{total_geral_bitola[bit]:.2f} kgf"
    )

print("\nTOTAL POR AÇO:")

for aco in ("50A", "60A"):
    if aco in total_geral_aco:
        print(
            f"  {aco}: "
            f"{total_geral_aco[aco]:.2f} kgf"
        )


from openpyxl import Workbook

# EXPORTA PARA EXCEL
from openpyxl import load_workbook

# LINHAS DOS PAVIMENTOS DENTRO DO EXCEL

LINHAS_PAVIMENTOS = {
    # Estrutura base (BLOCOS E PILARES,ENTRETANTO SE HOUVER FERRO NO NIVEL FUNDAÇÃO PODE MUDAR)
    "Fundação": 27,    
    "Pilares": 26,

    # Subsolos
    "Subsolo-2": 25,
    "Subsolo-1": 24,

    # Acesso
    "Patamar":40,
    "Rampa": 23,
    "Calçada": 22,
    "Térreo": 21,

    # Garagens
    "GAR1": 20,
    "GAR2": 19,

    # Pavimentos
    "Lazer": 18,
    "Tipo-1": 17,
    "PAV 5": 16,
    "Tipo-6-9": 34,
    "Tipo-10-15": 35,
    "Tipo-16-19": 36,
    "Tipo-20-25": 37,
    "Tipo-26-27": 38,
    "TIPO": 15,

    # Coberturas especiais
    "Duplex-inferior": 14,
    "Patamar-Superior": 13,
    "Duplex-superior": 12,
    "Cobertura": 11,
    "Salão-festas": 39,
    "Telhado": 10,
    "MOTOR": 9,
    "Caixa": 8,
    "Cinta": 7,
}

# COLUNAS DA PLANILHA

COLUNAS = {

    "Vigas": {
        5: "C",
        6.3: "D",
        8: "E",
        10: "F",
        12.5: "G",
        16: "H",
        20: "I",
    },

    "Lajes": {
        5: "J",
        6.3: "K",
        8: "L",
        10: "M",
        12.5: "N",
        16: "O",
        20: "P",
    },

    "Escadas": {
        5: "J",
        6.3: "K",
        8: "L",
        10: "M",
        12.5: "N",
        16: "O",
        20: "P",
    },

    "Pilares": {
        5: "R",
        6.3: "S",
        8: "T",
        10: "U",
        12.5: "V",
        16: "W",
        20: "X",
        25: "Y",
    },

    "Blocos": {
        5: "R",
        6.3: "S",
        8: "T",
        10: "U",
        12.5: "V",
        16: "W",
        20: "X",
        25: "Y",
    }
}


# PLANILHA MODELO (O engenheiro deve ter uma planilha base do excel e ajustar as linhas e colunas dos elementos e pavimentos de acordo)

MODELO = r"c:\caminho\da\pasta\planilha-ferros.xlsx" #caminho da pasta aonde está a planilha modelo .xlsx

wb = load_workbook(MODELO)
ws = wb.active

# COLOCAR NOME DO PAV DO LADO

for pavimento in dados.keys():

    linha = LINHAS_PAVIMENTOS.get(pavimento)

    if linha is not None:
        ws[f"B{linha}"] = pavimento

# PREENCHE A PLANILHA

for pav, elementos in dados.items():

    linha = LINHAS_PAVIMENTOS.get(pav)

    if linha is None:
        print(f" Pavimento sem linha definida: {pav}")
        continue

    for elem, bitolas in elementos.items():

        mapa_colunas = COLUNAS.get(elem)

        if mapa_colunas is None:
            print(f" Elemento sem coluna definida: {elem}")
            continue

        for bit, peso in bitolas.items():

            coluna = mapa_colunas.get(bit)

            if coluna is None:
                print(
                    f" Sem coluna para "
                    f"{pav} / {elem} / {bit} mm"
                )
                continue

            celula = f"{coluna}{linha}"

            valor_atual = ws[celula].value

            if valor_atual is None:
                valor_atual = 0

            try:
                ws[celula] = round(float(valor_atual) + peso, 2)

            except Exception as e:
                print(
                    f"Erro na célula {celula} | "
                    f"valor_atual={valor_atual!r} | "
                    f"tipo={type(valor_atual)} | "
                    f"peso={peso}"
                )
                raise

# TOTAL GERAL
ws["Z1"] = round(total_geral, 2)

# SALVA
saida = r"c:\caminho\da\pasta\resumos.xlsx" #caminho da pasta que será salvada

wb.save(saida)

print(f"\n Planilha gerada com sucesso:")
print(saida)